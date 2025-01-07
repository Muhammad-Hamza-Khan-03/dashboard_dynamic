import datetime
from mixtral import ImprovedMermaidGenerator
import os
import json
import re
import io
import logging
import tempfile
from typing import Dict, Any, List, Optional
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import plotly.express as px
import sqlite3
from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from pdfminer.high_level import extract_text
from docx import Document
from striprtf.striprtf import rtf_to_text
from dotenv import load_dotenv, find_dotenv
from langchain.prompts import ChatPromptTemplate
from langchain_core.messages import HumanMessage, AIMessage, SystemMessage
from langchain_groq import ChatGroq
from langchain.memory import ConversationBufferMemory
import PyPDF2
import ast
import traceback

class DataExplorerProAPI:
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.mermaid_generator = ImprovedMermaidGenerator()
        self._initialize_environment()
        self._setup_app()
        self._setup_model_and_memory()

    def _initialize_environment(self):
        dotenv_path = find_dotenv()
        load_dotenv(dotenv_path)
        self.GROQ_API_KEY = os.getenv("GROQ_API_KEY")

    def _setup_app(self):
        self.app = Flask(__name__)
        CORS(self.app)
        self._setup_routes()

    def _setup_model_and_memory(self):
        self.model = ChatGroq(api_key=self.GROQ_API_KEY, model="llama-3.3-70b-specdec")
        self.memory = ConversationBufferMemory(return_messages=True)

    def _setup_routes(self):
        routes = [
            ('/', 'index', self.index, ['GET']),
            ('/upload', 'upload_file', self.upload_file, ['POST']),
            ('/query', 'query_file', self.query_file, ['POST']),
            ('/data_stats', 'get_data_stats', self.get_data_stats, ['GET']),
            ('/download_cleaned_data', 'download_cleaned_data', self.download_cleaned_data, ['GET']),
            ('/download_csv', 'download_csv', self.download_csv, ['GET']),
            ('/export_graph', 'export_graph', self.export_graph, ['POST'])
        ]
        for route in routes:
            self.app.add_url_rule(route[0], route[1], route[2], methods=route[3])

    def index(self):
        return jsonify({'message': 'Welcome to the Data Explorer Pro API'})

    def upload_file(self):
        if 'file' not in request.files:
            return jsonify({'message': 'No file part'}), 400

        file = request.files['file']
        if file.filename == '':
            return jsonify({'message': 'No selected file'}), 400

        try:
            file_extension = os.path.splitext(file.filename)[1].lower()
            if file_extension in ['.csv', '.tsv']:
                return self._handle_csv_tsv_upload(file, file_extension)
            elif file_extension in ['.xls', '.xlsx']:
                return self._handle_excel_upload(file)
            elif file_extension in ['.sqlite', '.db']:
                return self._handle_sqlite_upload(file)
            elif file_extension in ['.txt', '.pdf', '.rtf', '.doc', '.docx']:
                return self._handle_document_upload(file, file_extension)
            else:
                return jsonify({'message': 'Unsupported file type'}), 400
        except Exception as e:
            logging.error(f"Error processing file: {e}\n{traceback.format_exc()}")
            return jsonify({'message': f'Error processing file: {str(e)}'}), 500

    def _handle_csv_tsv_upload(self, file, file_extension):
        df = pd.read_csv(file, sep=',' if file_extension == '.csv' else '\t')
        df.to_pickle('dataframe.pkl')
        df = df.fillna("NaN")
        preview = {
            'columns': df.columns.tolist(),
            'rows': df.head(1).values.tolist()
        }
        return jsonify({'message': f'{file_extension[1:].upper()} file uploaded successfully', 'preview': preview})

    def _handle_excel_upload(self, file):
        df = pd.read_excel(file)
        df.to_pickle('dataframe.pkl')
        df = df.fillna("NaN")
        preview = {
            'columns': df.columns.tolist(),
            'rows': df.head(1).values.tolist()
        }
        return jsonify({'message': 'Excel file uploaded successfully', 'preview': preview})

    def _handle_sqlite_upload(self, file):
        temp_dir = tempfile.mkdtemp()
        temp_path = os.path.join(temp_dir, file.filename)
        file.save(temp_path)

        conn = sqlite3.connect(temp_path)
        cursor = conn.cursor()
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table';")
        tables = [table[0] for table in cursor.fetchall()]

        self.db_info = {}
        for table in tables:
            cursor.execute(f"PRAGMA table_info({table});")
            columns = cursor.fetchall()
            column_names = [col[1] for col in columns]
            cursor.execute(f"SELECT * FROM {table} LIMIT 5;")
            rows = cursor.fetchall()
            self.db_info[table] = {
                'columns': column_names,
                'rows': rows
            }

        conn.close()
        os.remove(temp_path)
        os.rmdir(temp_dir)

        return jsonify({'message': 'SQLite file uploaded successfully', 'db_info': self.db_info})

    def _handle_document_upload(self, file, file_extension):
        try:
            if file_extension == '.txt':
                self.document_content = file.read().decode('utf-8')
            elif file_extension == '.pdf':
                pdf_reader = PyPDF2.PdfReader(file)
                self.document_content = "".join(page.extract_text() for page in pdf_reader.pages)
            elif file_extension == '.rtf':
                rtf_text = file.read().decode('utf-8')
                self.document_content = rtf_to_text(rtf_text)
            elif file_extension in ['.doc', '.docx']:
                doc = Document(file)
                self.document_content = "\n".join([paragraph.text for paragraph in doc.paragraphs])
            else:
                self.document_content = ""
        except Exception as e:
            logging.error(f"Error processing document: {e}\n{traceback.format_exc()}")
            return jsonify({'message': f'Error processing document: {str(e)}'}), 500

        preview = self.document_content[:1000]
        return jsonify({'message': f'{file_extension[1:].upper()} file uploaded successfully', 'preview': preview})

    def handle_data_cleaning(self, user_input):
        if not os.path.exists('dataframe.pkl'):
            return jsonify({'error': 'No dataset has been uploaded for cleaning'}), 400

        df = pd.read_pickle('dataframe.pkl')
        original_df = df.copy()
        df_info = self._get_dataframe_info(df)
        
        markdown_table = df.head(10).to_markdown(index=False)
        
        data_cleaning_instructions = self._get_data_cleaning_instructions(df_info, markdown_table, user_input)

        system_message = SystemMessage(content=data_cleaning_instructions)
        messages = [system_message, HumanMessage(content=user_input)]

        try:
            response = self.model(messages)
            result_output = response.content

            code_blocks = re.findall(r'```python(.*?)```', result_output, re.DOTALL)
            for code_block in code_blocks:
                code_block = code_block.strip()
                logging.debug(f"Executing code block:\n{code_block}")
                if self.is_valid_python(code_block):
                    try:
                        exec(code_block, {'df': df, 'pd': pd, 'np': np})
                        logging.debug("Code block executed successfully.")
                    except SyntaxError as se:
                        logging.error(f"SyntaxError executing code block: {se}\nCode:\n{code_block}")
                        result_output += f"\n\nSyntax error executing code: {str(se)}"
                    except Exception as e:
                        logging.error(f"Error executing code block: {e}\nCode:\n{code_block}")
                        result_output += f"\n\nError executing code: {str(e)}"
                else:
                    logging.error(f"Invalid Python code provided:\n{code_block}")
                    result_output += "\n\nInvalid Python code provided."

            preview_data = df.head(10).to_dict('records')

            cleaned_markdown_table = df.head(10).to_markdown(index=False)

            result_output += self._get_dataset_changes_info(original_df, df, cleaned_markdown_table)

            suggestions = self.generate_suggested_prompts(user_input, df=df, agent_type='data_cleaning')
            response_json = jsonify({
                'output': result_output,
                'preview_data': preview_data,
                'total_rows': len(df),
                'suggestions': suggestions,
                'dataframe_changed': not original_df.equals(df)
            })
            
            self.cleaned_df = df
            
            logging.debug(f"Data cleaning response: {response_json.get_data(as_text=True)}")
            return response_json
        except Exception as e:
            logging.error(f"Error in data cleaning: {e}\n{traceback.format_exc()}")
            return jsonify({'output': f'Error: {str(e)}'}), 500

    def is_valid_python(self, code):
        """
        Validate Python code syntax using the ast module.
        """
        try:
            ast.parse(code)
            return True
        except SyntaxError as se:
            logging.error(f"Invalid Python code: {se}\nCode:\n{code}")
            return False

    def _get_dataframe_info(self, df):
        info = {
            'columns': df.columns.tolist(),
            'total_rows': len(df),
            'total_columns': len(df.columns)
        }
        return info

    def _get_data_cleaning_instructions(self, df_info, markdown_table, user_input):
        return f"""You are a data cleaning expert. Analyze the given dataset and provide recommendations for cleaning and preprocessing. Use the following information:

1. Dataset Overview:
- Columns: {', '.join(df_info['columns'])}
- Total rows: {df_info['total_rows']}
- Total columns: {df_info['total_columns']}

2. Sample Data (first 10 rows):
{markdown_table}

3. User Request:
{user_input}

Please provide the following:
1. An overview of data quality issues
2. Recommendations for handling missing values
3. Suggestions for data type conversions if necessary
4. Identification of potential outliers or anomalies
5. Proposals for feature engineering or data transformations
6. Any other relevant data cleaning steps

Use markdown formatting for better readability. Include Python code snippets wrapped in triple backticks for any cleaning operations you recommend.
After each cleaning operation, provide a brief explanation of what was done and why.
IMPORTANT: Do not include any code that saves or overwrites the dataframe. The system will handle dataset versioning automatically.
"""

    def _get_dataset_changes_info(self, original_df, df, cleaned_markdown_table):
        return f"\n\n### Dataset Changes\n" \
               f"Original rows: {len(original_df)}, Current rows: {len(df)}\n" \
               f"Original columns: {', '.join(original_df.columns)}\n" \
               f"Current columns: {', '.join(df.columns)}\n" \
               f"\n### Cleaned Data Preview (First 10 rows):\n{cleaned_markdown_table}\n"

    def download_cleaned_data(self):
        if not hasattr(self, 'cleaned_df'):
            return jsonify({'error': 'No cleaned dataset available'}), 400

        return self._send_dataframe_as_csv(self.cleaned_df, 'cleaned_data.csv')

    def handle_mermaid_diagram(self, user_input):
        try:
            # Log before generating diagram
            logging.info(f"Generating Mermaid diagram for input: {user_input}")
            
            result = self.mermaid_generator.generate_diagram(user_input)
            
            # Log the generated diagram code
            if result.get('mermaid'):
                logging.info("Generated Mermaid Code:")
                logging.info("\n" + result['mermaid'])
            else:
                logging.warning("No Mermaid code was generated")
                
            # Include the Mermaid code in the response for easier debugging
            response = {
                'output': result['output'],
                'mermaid': result['mermaid'],
                'suggestions': result['suggestions'],
                'debug': {
                    'mermaid_code': result['mermaid'],
                }
            }
            
            logging.debug(f"Full response: {response}")
            return jsonify(response)
            
        except Exception as e:
            logging.error(f"Error in Mermaid diagram generation: {e}\n{traceback.format_exc()}")
            return jsonify({'output': f'Error: {str(e)}', 'mermaid': None, 'suggestions': []})

    def query_file(self):
        user_input = request.json.get('question')
        agent_type = request.json.get('agent_type', 'data_visualization')
        custom_instructions = request.json.get('custom_instructions')

        if not user_input:
            return jsonify({'error': 'No question provided'}), 400

        if not os.path.exists('dataframe.pkl') and not hasattr(self, 'document_content') and 'file' not in request.files:
            return jsonify({'error': 'No file has been uploaded yet'}), 400

        if agent_type == 'research_assistant':
            return self._handle_research_assistant(user_input)
        elif agent_type == 'mermaid_diagram':
            return self.handle_mermaid_diagram(user_input)
        elif agent_type == 'data_cleaning':
            return self.handle_data_cleaning(user_input)
        elif agent_type == 'business_analytics':
            return self._handle_business_analytics(user_input)
        elif agent_type == 'sql' and hasattr(self, 'db_info') and self.db_info:
            return self._handle_sql_query(user_input)
        else:
            return self._handle_data_visualization(user_input, custom_instructions)

    def _handle_research_assistant(self, user_input):
        if not hasattr(self, 'document_content'):
            return jsonify({'error': 'No document has been uploaded for analysis'}), 400

        research_instructions = self._get_research_instructions(user_input)

        system_message = SystemMessage(content=research_instructions)
        messages = [system_message, HumanMessage(content=user_input)]

        try:
            response = self.model(messages)
            result_output = response.content
            suggestions = self.generate_suggested_prompts(user_input, agent_type='research_assistant')
            response_json = jsonify({'output': result_output, 'suggestions': suggestions})
            logging.debug(f"Research assistant response: {response_json.get_data(as_text=True)}")
            return response_json
        except Exception as e:
            logging.error(f"Error getting response from model: {e}\n{traceback.format_exc()}")
            return jsonify({'output': f'Error: {str(e)}'}), 500

    def _get_research_instructions(self, user_input):
        return f"""You are a research assistant analyzing a document. The user has the following question:

User Question: {user_input}

Please provide an informative answer based on the document's content. If you cannot answer the question based on the given information, please state that clearly.

Document Content (first 1000 characters):
{self.document_content[:1000]}

Important Notes:
- Base your response solely on the information provided in the document.
- If the answer is not in the given content, say so clearly.
- Use markdown formatting in your response for better readability.

Your response:"""

    def _handle_business_analytics(self, user_input):
        if not os.path.exists('dataframe.pkl'):
            return jsonify({'error': 'No dataset has been uploaded for analysis'}), 400

        df = pd.read_pickle('dataframe.pkl')
        df_info = self._get_dataframe_info(df)
        
        business_analytics_instructions = self._get_business_analytics_instructions(df_info, df.head().to_string(), user_input)

        system_message = SystemMessage(content=business_analytics_instructions)
        messages = [system_message, HumanMessage(content=user_input)]

        try:
            response = self.model(messages)
            result_output = response.content

            code_blocks = re.findall(r'```python(.*?)```', result_output, re.DOTALL)
            for code_block in code_blocks:
                code_block = code_block.strip()
                logging.debug(f"Executing code block:\n{code_block}")
                if self.is_valid_python(code_block):
                    try:
                        exec(code_block, {'df': df, 'pd': pd, 'np': np})
                        logging.debug("Code block executed successfully.")
                    except SyntaxError as se:
                        logging.error(f"SyntaxError executing code block: {se}\nCode:\n{code_block}")
                        result_output += f"\n\nSyntax error executing code: {str(se)}"
                    except Exception as e:
                        logging.error(f"Error executing code block: {e}\nCode:\n{code_block}")
                        result_output += f"\n\nError executing code: {str(e)}"
                else:
                    logging.error(f"Invalid Python code provided:\n{code_block}")
                    result_output += "\n\nInvalid Python code provided."

            suggestions = self.generate_suggested_prompts(user_input, df=df, agent_type='business_analytics')
            response_json = jsonify({'output': result_output, 'suggestions': suggestions})
            logging.debug(f"Business analytics response: {response_json.get_data(as_text=True)}")
            return response_json
        except Exception as e:
            logging.error(f"Error getting response from model: {e}\n{traceback.format_exc()}")
            return jsonify({'output': f'Error: {str(e)}'}), 500

    def _get_business_analytics_instructions(self, df_info, sample_data, user_input):
        return f"""You are a business analytics expert. Analyze the given dataset and provide key trends, insights, and focus on profit/revenue trends. Use the following information:
1. Dataset Overview:
- Columns: {', '.join(df_info['columns'])}
- Total rows: {df_info['total_rows']}
- Total columns: {df_info['total_columns']}

2. Sample Data (first 5 rows):
{sample_data}

3. User Request:
{user_input}

Please provide the following:
1. An overview of the dataset
2. Key trends and insights
3. Profit/revenue analysis (if applicable)
4. Recommendations based on the data
5. Potential areas for further investigation

Use markdown formatting for better readability. Include relevant statistics and percentages where appropriate.
If you need to perform any calculations, use Python code snippets wrapped in triple backticks.
"""

    def _handle_sql_query(self, user_input):
        if not hasattr(self, 'db_info'):
            return jsonify({'error': 'No database information available'}), 400

        db_info_str = json.dumps(self.db_info, indent=4)
        instructions = f"""You are a SQL expert. The dataset is stored in the SQLite database. Please write a SQL query based on the user's request and execute them.
        
The database contains the following tables and sample data:

{db_info_str}

User Request:
{user_input}

Please generate SQL queries based on the user's instructions. Explain the purpose of each query and what insights it aims to derive from the data.
"""

        system_message = SystemMessage(content=instructions)
        history = self.memory.load_memory_variables({})
        messages = [system_message] + history.get("history", []) + [HumanMessage(content=user_input)]

        try:
            response = self.model(messages)
        except Exception as e:
            logging.error(f"Error getting response from model: {e}\n{traceback.format_exc()}")
            return jsonify({'output': f'Error: {str(e)}'}), 500

        self.memory.save_context({"input": user_input}, {"output": response.content})
        result_output = response.content
        
        # Extract SQL queries from the response
        sql_queries = re.findall(r'```sql\n(.*?)\n```', result_output, re.DOTALL)
        
        # Execute queries and format results
        query_results = []
        for query in sql_queries:
            query = query.strip()
            if not query.endswith(';'):
                query += ';'
            try:
                results = self._execute_sql_query(query)
                formatted_results = [self._format_sql_result(result) for result in results]
                query_results.append({
                    'query': query,
                    'results': formatted_results
                })
            except Exception as e:
                logging.error(f"Error executing SQL query: {e}\nQuery: {query}\n{traceback.format_exc()}")
                query_results.append({
                    'query': query,
                    'error': str(e)
                })
        combined_output = result_output + "\n\n## Query Results\n\n"
        for i, result in enumerate(query_results, 1):
            combined_output += f"### Query {i}\n"
            combined_output += f"```sql\n{result['query']}\n```\n\n"
            if 'results' in result:
                for j, res in enumerate(result['results'], 1):
                    combined_output += f"#### Result {j}:\n" + res + "\n\n"
            else:
                combined_output += f"#### Error: {result['error']}\n\n"

        suggestions = self.generate_suggested_prompts(user_input, db_info=self.db_info, agent_type='sql')
        
        response_json = {'output': combined_output, 'suggestions': suggestions}
        logging.debug(f"SQL query response: {response_json}")
        return jsonify(response_json)

    def _execute_sql_query(self, query):        
        conn = sqlite3.connect(':memory:')
        cursor = conn.cursor()

        for table_name, table_info in self.db_info.items():
            if table_name.lower() != 'sqlite_sequence':  
                columns = ', '.join([f"{col} TEXT" for col in table_info['columns']])  # Assuming all columns are TEXT for simplicity
                create_table_query = f"CREATE TABLE {table_name} ({columns})"
                try:
                    cursor.execute(create_table_query)
                except sqlite3.Error as e:
                    logging.error(f"Error creating table {table_name}: {e}\n{traceback.format_exc()}")

                # Insert sample data
                for row in table_info['rows']:
                    placeholders = ', '.join(['?' for _ in row])
                    insert_query = f"INSERT INTO {table_name} VALUES ({placeholders})"
                    try:
                        cursor.execute(insert_query, row)
                    except sqlite3.Error as e:
                        logging.error(f"Error inserting data into table {table_name}: {e}\n{traceback.format_exc()}")

        statements = query.split(';')
        results = []

        for statement in statements:
            statement = statement.strip()
            if statement:
                try:
                    cursor.execute(statement)
                    if cursor.description:  # If the query returns results
                        fetched = cursor.fetchall()
                        column_names = [description[0] for description in cursor.description]
                        results.append({'columns': column_names, 'rows': fetched})
                    else:
                        results.append({'columns': [], 'rows': []})
                except sqlite3.Error as e:
                    logging.error(f"SQLite error: {e}\nQuery: {statement}\n{traceback.format_exc()}")
                    results.append({'error': str(e)})

        conn.close()
        return results

    def _format_sql_result(self, result):
        if 'error' in result:
            return f"Error: {result['error']}"

        columns = result.get('columns', [])
        rows = result.get('rows', [])

        if not columns and not rows:
            return "No results returned."

        col_widths = [len(str(col)) for col in columns]
        for row in rows:
            for i, cell in enumerate(row):
                col_widths[i] = max(col_widths[i], len(str(cell)))

        header = "| " + " | ".join(str(col).ljust(col_widths[i]) for i, col in enumerate(columns)) + " |"
        separator = "|-" + "-|-".join("-" * width for width in col_widths) + "-|"

        formatted_rows = []
        for row in rows:
            formatted_row = "| " + " | ".join(str(cell).ljust(col_widths[i]) for i, cell in enumerate(row)) + " |"
            formatted_rows.append(formatted_row)
        table = "\n".join([header, separator] + formatted_rows)
        return table

    def _handle_data_visualization(self, user_input, custom_instructions):
        if not os.path.exists('dataframe.pkl'):
            return jsonify({'error': 'No dataset has been uploaded for visualization'}), 400

        df = pd.read_pickle('dataframe.pkl')
        df_5_rows = df.head(5).fillna("NaN")
        column_names = ', '.join(df.columns)
        csv_string = df_5_rows.to_string(index=False)
        total_rows = len(df)
        total_columns = len(df.columns)

        instructions = custom_instructions or self._get_default_visualization_instructions(
            total_rows, total_columns, column_names, csv_string, user_input
        )

        system_message = SystemMessage(content=instructions)
        history = self.memory.load_memory_variables({})
        messages = [system_message] + history.get("history", []) + [HumanMessage(content=user_input)]

        try:
            response = self.model(messages)
        except Exception as e:
            logging.error(f"Error getting response from model: {e}\n{traceback.format_exc()}")
            return jsonify({'graph': None, 'output': f'Error: {str(e)}', 'suggestions': []})

        self.memory.save_context({"input": user_input}, {"output": response.content})
        result_output = response.content
        suggestions = self.generate_suggested_prompts(user_input, df=df, agent_type='data_visualization')
        
        response_json = {'graph': None, 'output': result_output, 'suggestions': suggestions}

        code_block_match = re.search(r'```(?:Pp|python)?(.*?)```', result_output, re.DOTALL)
        if code_block_match:
            code_block = code_block_match.group(1).strip()
            cleaned_code = re.sub(r'(?m)^\s*fig\.show\(\)\s*$', '', code_block)
            
            printed_output = io.StringIO()
            try:
                if self.is_valid_python(cleaned_code):
                    exec(cleaned_code, {'pd': pd, 'px': px, 'df': df, 'print': lambda *args: printed_output.write(" ".join(map(str, args)) + "\n")})
                    fig = self.get_fig_from_code(cleaned_code, df)
                    if isinstance(fig, str):
                        response_json['output'] += fig
                    elif fig:
                        fig_json = fig.to_json()
                        response_json['graph'] = fig_json
                        response_json['output'] += result_output
                    else:
                        response_json['output'] += 'No figure was generated.'
                else:
                    response_json['output'] += 'Invalid Python code provided.'
            except SyntaxError as e:
                logging.error(f"Syntax error in code: {e}\nCode:\n{cleaned_code}\n{traceback.format_exc()}")
                response_json['output'] += f'Syntax error in code: {str(e)}'
            except Exception as e:
                logging.error(f"Error executing code: {e}\nCode:\n{cleaned_code}\n{traceback.format_exc()}")
                response_json['output'] += f'Error executing code: {str(e)}'

        logging.debug(f"Query response: {response_json}")
        return jsonify(response_json)

    def _get_default_visualization_instructions(self, total_rows, total_columns, column_names, csv_string, user_input):
        return f"""You are a data visualization assistant using the Plotly library. The dataset is stored in the `df` variable. You will be provided with information about the dataset and instructions to create a graph. Please follow these guidelines:

1. Dataset Overview:
- Total rows: {total_rows}
- Total columns: {total_columns}
- Column names: {column_names}
    
2. Sample Data:
Here are the first 5 rows of the dataset:
{csv_string}

3. Important Notes:
- The full dataset is available in the `df` variable.
- Use `df.dtypes` to check column data types if needed.
- Handle potential missing values appropriately.
- Consider the entire dataset when making visualizations or analyses.

4. User Request:
{user_input}
    
5. Markdown Formatting:
- Use Markdown formatting in your responses.
- Wrap SQL queries and results in triple backticks (```).
- Use # for headers, ## for subheaders, etc.
- Use * or - for bullet points.
- Use **text** for bold and *text* for italics.
- Format result tables using | for columns and - for separators.

Please follow the user's instructions to generate the appropriate graph or analysis. Ensure your code is efficient and can handle the full dataset. If the user's request is unclear or could be interpreted in multiple ways, ask for clarification.

After executing the code, explain your approach and any insights gained from the visualization or analysis. Always show the number of rows using markdown table. If there are any limitations or potential issues with the requested visualization given the nature of the data, mention them."""

    def get_data_stats(self):
        if not os.path.exists('dataframe.pkl'):
            return jsonify({'error': 'No CSV file has been uploaded yet'}), 400

        df = pd.read_pickle('dataframe.pkl')

        df = df.fillna("NaN")
        stats = {
            'describe': df.describe(include='all').to_dict(),
            'correlation': df.corr(numeric_only=True).to_dict()
        }
        return jsonify(stats)

    def download_csv(self):
        if not os.path.exists('dataframe.pkl'):
            return jsonify({'error': 'No CSV file has been uploaded yet'}), 400

        df = pd.read_pickle('dataframe.pkl')
        return self._send_dataframe_as_csv(df, 'data.csv')

    def _send_dataframe_as_csv(self, df, filename):
        csv_buffer = io.StringIO()
        df.to_csv(csv_buffer, index=False)
        csv_buffer.seek(0)
        
        return send_file(
            io.BytesIO(csv_buffer.getvalue().encode()),
            mimetype='text/csv',
            as_attachment=True,
            download_name=filename
        )

    def export_graph(self):
        try:
            logging.debug("Received request to export graph")

            graph_json = request.json.get('graph')
            logging.debug(f"Received graph JSON: {graph_json}")

            if not graph_json:
                return jsonify({'error': 'No graph data provided'}), 400

            try:
                graph_data = json.loads(graph_json)
                logging.debug(f"Parsed graph data: {graph_data}")
            except json.JSONDecodeError as e:
                logging.error(f"Error decoding graph JSON: {e}\n{traceback.format_exc()}")
                return jsonify({'error': 'Invalid graph JSON data'}), 400

            return jsonify(graph_data)

        except Exception as e:
            logging.error(f"Error exporting graph: {e}\n{traceback.format_exc()}")
            return jsonify({'error': 'An error occurred while exporting the graph'}), 500

    def get_fig_from_code(self, code, df):
        local_variable = {}
        try:
            logging.debug(f"Executing code: {code}")
            exec(code, {'pd': pd, 'px': px, 'df': df}, local_variable)
            return local_variable.get('fig', None)
        except SyntaxError as e:
            logging.error(f"Syntax error in code: {e}\nCode:\n{code}\n{traceback.format_exc()}")
            return f"Syntax error in code: {e}"
        except Exception as e:
            logging.error(f"Error executing code: {e}\nCode:\n{code}\n{traceback.format_exc()}")
            return f"Error executing code: {e}"

    def generate_suggested_prompts(self, user_input, df=None, db_info=None, agent_type='data_visualization'):
        prompt_generators = {
            'research_assistant': self._generate_research_assistant_prompt,
            'sql': self._generate_sql_prompt,
            'data_cleaning': self._generate_data_cleaning_prompt,
            'business_analytics': self._generate_business_analytics_prompt,
            'data_visualization': self._generate_data_visualization_prompt
        }

        prompt_generator = prompt_generators.get(agent_type, self._generate_data_visualization_prompt)
        prompt = prompt_generator(user_input, df, db_info)

        logging.debug(f"Prompt sent to model: {prompt}")

        try:
            response = self.model([
                SystemMessage(content="You are a helpful data analysis assistant."),
                HumanMessage(content=prompt)
            ])
            logging.debug(f"Raw response from model: {response.content}")
            
            suggestions = response.content.split('\n')
            suggestions = [s.strip() for s in suggestions if s.strip().endswith('?')]
            
            logging.debug(f"Filtered suggestions: {suggestions}")
            return suggestions[:3]
        except Exception as e:
            logging.error(f"Error generating suggestions: {e}\n{traceback.format_exc()}")
            return ["Error generating suggestions. Please try again."]

    def _generate_research_assistant_prompt(self, user_input, df=None, db_info=None):
        return f"""Based on the following user input and the context of document analysis, suggest 3 follow-up questions that the user might find interesting or useful for further exploration. Make sure the suggestions are relevant to both the user's query and the nature of document analysis.
User input: {user_input}

IMPORTANT: Provide ONLY the questions, one per line. Do not include any explanations, numbering, or additional text.

Suggested questions:"""

    def _generate_sql_prompt(self, user_input, df=None, db_info=None):
        db_info_str = json.dumps(db_info, indent=4)
        return f"""Analyze the user's input and the provided database schema. Verify if the necessary columns exist and then generate 3 insightful SQL-oriented follow-up questions that:
1. Dig deeper into the user's initial query
2. Explore related aspects of the data
3. Uncover potential trends or patterns

Ensure each question:
- Is directly executable as a SQL query
- Utilizes appropriate tables and columns from the schema
- Incorporates relevant SQL functions or operations
- Avoids redundancy with the original query

Format: Present only the questions, one per line, without numbering or explanation. Also show any row/s or table using the markdown table.

User input: {user_input}

Database schema:
{db_info_str}

SQL follow-up questions:"""

    def _generate_data_cleaning_prompt(self, user_input, df=None, db_info=None):
        if df is not None:
            df_info = self._get_dataframe_info(df)
            missing_values = df.isnull().sum().to_dict()
            missing_values_str = json.dumps(missing_values, indent=4)
        else:
            df_info = "No dataframe available"
            missing_values_str = "No missing value information available"

        return f"""Based on the following user input and the context of data cleaning, suggest 3 follow-up questions that the user might find interesting or useful for further data preparation. Focus on data quality improvements, handling missing values, data transformations, and feature engineering.

User input: {user_input}

Dataset information:
Columns: {', '.join(df_info['columns']) if isinstance(df_info, dict) else df_info}
Total rows: {df_info['total_rows'] if isinstance(df_info, dict) else ''}
Total columns: {df_info['total_columns'] if isinstance(df_info, dict) else ''}

Missing values:
{missing_values_str}

IMPORTANT: Provide ONLY the questions, one per line. Do not include any explanations, numbering, or additional text.

Suggested questions:"""

    def _generate_business_analytics_prompt(self, user_input, df=None, db_info=None):
        df_info = self._get_dataframe_info(df) if df is not None else "No dataframe available"
        return f"""Based on the following user input and the context of business analytics, suggest 3 follow-up questions that the user might find interesting or useful for further exploration. Focus on profit/revenue trends, market insights, and potential business strategies.

User input: {user_input}

Dataset information:
Columns: {', '.join(df_info['columns']) if isinstance(df_info, dict) else df_info}
Total rows: {df_info['total_rows'] if isinstance(df_info, dict) else ''}
Total columns: {df_info['total_columns'] if isinstance(df_info, dict) else ''}

IMPORTANT: Provide ONLY the questions, one per line. Do not include any explanations, numbering, or additional text.

Suggested questions:"""

    def _generate_data_visualization_prompt(self, user_input, df=None, db_info=None):
        if df is not None:
            df_info = self._get_dataframe_info(df)
        else:
            df_info = "No dataframe available"
        return f"""Based on the following user input and dataframe information, suggest 3 follow-up questions that the user might find interesting or useful for further data exploration. Make sure the suggestions are relevant to both the user's query and the available data. 

IMPORTANT: Provide ONLY the questions, one per line. Do not include any explanations, numbering, or additional text.

User input: {user_input}

Dataframe information:
Columns: {', '.join(df_info['columns']) if isinstance(df_info, dict) else df_info}
Total rows: {df_info['total_rows'] if isinstance(df_info, dict) else ''}
Total columns: {df_info['total_columns'] if isinstance(df_info, dict) else ''}

Suggested questions:"""

    def run(self):
        self.app.run(debug=False, port=7000)

if __name__ == '__main__':
    logging.basicConfig(level=logging.DEBUG,
                        format='%(asctime)s %(levelname)s %(message)s',
                        handlers=[
                            logging.FileHandler("debug.log"),
                            logging.StreamHandler()
                        ])
    api = DataExplorerProAPI()
    api.run()
